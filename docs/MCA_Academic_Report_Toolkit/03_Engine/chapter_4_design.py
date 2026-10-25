"""
Master Report Builder — Chapter 4: System Design (Wi-Fi Sense Architecture)
MCA Academic Report Toolkit — Amal Jyothi College of Engineering Autonomous
Department of Computer Applications | Course: 24MCAR295
"""

import os
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from .styles import (
        add_chapter_heading, add_heading_1, add_heading_2, add_heading_3,
        add_body_p, add_bullet_p, add_styled_table, add_figure_image, add_code_block
    )
except ImportError:
    from styles import (
        add_chapter_heading, add_heading_1, add_heading_2, add_heading_3,
        add_body_p, add_bullet_p, add_styled_table, add_figure_image, add_code_block
    )
try:
    from .schemas import add_all_database_schemas
except ImportError:
    from schemas import add_all_database_schemas

def add_chapter_4_design(doc, cfg=None, diagrams_dir=None):
    if diagrams_dir is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        diagrams_dir = os.path.join(base_dir, "01_Inputs", "diagrams")
    if cfg is None:
        cfg = {}

    pj = cfg.get("project", {})
    project_title = pj.get("title", "Wi-Fi Sense: AI-Powered Indoor Monitoring and Facility Intelligence Platform Using Wi-Fi CSI")
    project_short = pj.get("short_name", "Wi-Fi Sense")

    # =========================================================================
    # CHAPTER 4: SYSTEM DESIGN
    # =========================================================================
    add_chapter_heading(doc, 4, "System Design")

    add_heading_1(doc, "4.1 INTRODUCTION")
    add_body_p(
        doc,
        f"The design phase is a crucial stage in the development of {project_short}, an AI-assisted indoor monitoring and "
        "facility intelligence platform designed to centralize contactless sensing, human activity classification, "
        "emergency fall detection, and space utilization analytics. A well-structured system design establishes the foundation "
        "for reliable operation, maintainability, strict clinical privacy, and horizontal scalability. This phase involves defining "
        "the physical RF edge architecture, database structure, graphical user interface, and integration of various micro-components."
    )
    add_body_p(
        doc,
        "The system design translates the requirements identified during the system study and feasibility analysis into a structured "
        "technical blueprint for implementation. It defines how the ESP32-S3 IoT transceivers capture 64-subcarrier OFDM channel states, "
        "how the host serial gateway daemon buffers and transmits telemetry, how the FastAPI ASGI backend sanitizes phase noise and evaluates "
        "activity states, and how real-time WebSocket streams push life-critical alerts to attending nurse consoles. "
        "By prioritizing usability, security, and extensibility, the architecture is designed to simplify clinical monitoring workflows "
        "while providing a solid foundation for future research enhancements."
    )

    add_heading_1(doc, "4.2 UML DIAGRAM")
    add_body_p(
        doc,
        "Unified Modeling Language (UML) is a standard visual modeling language used to specify, visualize, design, "
        "and document the structure and behavior of software systems. It provides a common set of diagrams that help engineers "
        "and stakeholders understand system components, interactions, workflows, and deployment architecture."
    )
    add_body_p(
        doc,
        f"In {project_short}, nine standardized UML diagrams are used to represent the system architecture, user interactions, "
        "internal components, and operational workflows. These diagrams provide a clear understanding of how clinical caregivers, "
        "facility managers, corporate staff, and family members interact with the platform, and how its distributed components "
        "communicate to support real-time contactless sensing and emergency dispatch."
    )

    # 4.2.1 USE CASE DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.1 Use Case Diagram")
    add_body_p(
        doc,
        "A use case diagram is a graphical representation that illustrates how users and external entities interact with a system. "
        "It helps identify, outline, and organize the system's functional requirements from an actor's perspective. "
        "Built using UML standards, use case diagrams define system boundaries, actors, use cases, and relationships."
    )
    add_body_p(
        doc,
        "Figure 4.1 shows the use cases of Wi-Fi Sense across its administrative and clinical access levels. "
        "Primary actors include the System Administrator, Organization Administrator, Facility Manager, Caregiver/Staff, "
        "Corporate Facility Staff, Family Member/Doctor, and the physical ESP32-S3 CSI Sensing Node."
    )
    add_bullet_p(doc, "Administration: Manage Organizations, Manage Users & Roles, Provision Facilities.", "Administration")
    add_bullet_p(doc, "Facility Management: Manage Buildings, Floors, Rooms, and Sensing Devices.", "Facility Management")
    add_bullet_p(doc, "Sensing & AI: Collect CSI Data, Preprocess CSI Data, Detect Presence, Recognize Activity, Detect Potential Fall.", "Sensing & AI")
    add_bullet_p(doc, "Monitoring: View Central Monitoring Dashboard, View Room Occupancy, Monitor Resident Presence, Monitor Activity.", "Monitoring")
    add_bullet_p(doc, "Alert Management: Generate Alert, Send Notification, View Alert History, Acknowledge Alert, Resolve Alert.", "Alert Management")
    add_bullet_p(doc, "Analytics: View Occupancy Analytics, View Room Utilization, View Historical Activity.", "Analytics")
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_01_use_case.png"),
        "Figure 4.1: Use case diagram of Wi-Fi Sense",
        width_inches=5.8
    )

    # 4.2.2 SEQUENCE DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.2 Sequence Diagram")
    add_body_p(
        doc,
        "A sequence diagram is a type of interaction diagram that depicts the chronological flow of messages between system components. "
        "It helps visualize object interactions, understand system requirements, and detect potential concurrency bottlenecks. "
        "Key elements include actors, lifelines, messages, and execution activations."
    )
    add_body_p(
        doc,
        "Figure 4.2 shows the end-to-end emergency fall detection and dispatch interaction sequence. When a resident falls, "
        "the ESP32-S3 node captures raw CSI data and transmits it to the CSI Processing Service. The processed features are passed "
        "to the ML Classification Service, which identifies a high-confidence fall. The backend creates an Alert record, "
        "and the Alert Engine broadcasts an instant notification to the Caregiver Console. The caregiver acknowledges the alert, "
        "investigates the room, and records the resolution."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_02_sequence.png"),
        "Figure 4.2: Sequence diagram for fall detection and caregiver dispatch in Wi-Fi Sense",
        width_inches=5.8
    )

    # 4.2.3 STATE CHART DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.3 State Chart Diagram")
    add_body_p(
        doc,
        "A State Chart Diagram is a UML behavioral diagram that represents the various states an entity can exist in and "
        "how it transitions between them in response to internal and external events. Key components include the Initial State, "
        "States, Transitions, Events, Actions, and Final States."
    )
    add_body_p(
        doc,
        "Figure 4.3 shows the lifecycle of an emergency alert record in Wi-Fi Sense. The states follow the alert status field: "
        "Detection, Created, Notified, Acknowledged, and Resolved. Transitions between states require explicit caregiver "
        "actions and log responder identity for clinical accountability."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_03_state.png"),
        "Figure 4.3: State chart diagram of an emergency alert record",
        width_inches=5.8
    )

    # 4.2.4 ACTIVITY DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.4 Activity Diagram")
    add_body_p(
        doc,
        "An activity diagram illustrates the procedural flow of control from activity to activity within a system workflow. "
        "It models sequential steps, decision branching points, and concurrent parallel activities using standardized symbols:"
    )
    add_bullet_p(doc, "Initial node: Black circle representing workflow starting point.", "Initial Node")
    add_bullet_p(doc, "Activity: Rounded rectangle containing the operational task.", "Activity")
    add_bullet_p(doc, "Decision node: Diamond shape representing conditional branching.", "Decision Node")
    add_bullet_p(doc, "Final node: Encircled black circle representing workflow completion.", "Final Node")
    add_body_p(
        doc,
        "Figure 4.4 shows the activity workflow for continuous CSI packet processing, feature extraction, activity classification, "
        "and emergency escalation."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_04_activity.png"),
        "Figure 4.4: Activity diagram for CSI ingestion and fall incident workflow",
        width_inches=5.8
    )

    # 4.2.5 CLASS DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.5 Class Diagram")
    add_body_p(
        doc,
        "A Class Diagram is a key UML component used in object-oriented modeling to represent the static structure of an application. "
        "It defines classes, attributes, methods, and relationships such as association, aggregation, composition, and inheritance."
    )
    add_body_p(
        doc,
        "Figure 4.5 shows the persistent domain classes of Wi-Fi Sense matching the relational database schema: "
        "Organization, Building, Floor, Room, User, Role, SensingDevice, SensingEvent, ActivityType, Alert, and Notification."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_05_class.png"),
        "Figure 4.5: Class diagram of Wi-Fi Sense domain entities and relations",
        width_inches=5.8
    )

    # 4.2.6 OBJECT DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.6 Object Diagram")
    add_body_p(
        doc,
        "Object diagrams represent a snapshot of the system's runtime state at a particular point in time, showing concrete instances "
        "of classes and their attribute values."
    )
    add_body_p(
        doc,
        "Figure 4.6 shows an object diagram with instantiated entities during an active fall event: Organization (St. Peter's Elder Care Home), "
        "Room (Room 204), Device (ESP32-S3 Node), Resident (Annamma Joseph), Alert (Fall_Detected, Severity CRITICAL), and Caregiver (Sr. Mary Sebastian)."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_06_object.png"),
        "Figure 4.6: Object diagram showing runtime instance state during active fall event",
        width_inches=5.8
    )

    # 4.2.7 COMPONENT DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.7 Component Diagram")
    add_body_p(
        doc,
        "A component diagram illustrates how software components are organized and wired together to form larger subsystems. "
        "Figure 4.7 illustrates the modular subsystems of Wi-Fi Sense: the React Web Client, FastAPI API Gateway, Authentication Service, "
        "CSI Signal Processing Service, WiFall Inference Engine, Alert Management Service, and Supabase PostgreSQL Database."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_07_component.png"),
        "Figure 4.7: Component diagram of Wi-Fi Sense software architecture",
        width_inches=5.8
    )

    # 4.2.8 DEPLOYMENT DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.8 Deployment Diagram")
    add_body_p(
        doc,
        "A deployment diagram models the physical hardware nodes and runtime environments on which the software artifacts execute. "
        "Figure 4.8 shows the distributed deployment topology: the ESP32-S3 IoT sensing hardware in monitored rooms, the Host Gateway PC "
        "connected via USB COM port, the Render Linux Cloud container running the FastAPI backend, the Supabase Cloud PostgreSQL database, "
        "and client web browsers connected over secure HTTPS and WSS."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_08_deployment.png"),
        "Figure 4.8: Deployment diagram of Wi-Fi Sense hardware and cloud topology",
        width_inches=5.8
    )

    # 4.2.9 COLLABORATION DIAGRAM
    doc.add_page_break()
    add_heading_2(doc, "4.2.9 Collaboration Diagram")
    add_body_p(
        doc,
        "A collaboration (communication) diagram shows the structural relationships between cooperating objects, emphasizing the message links "
        "connecting them. Figure 4.9 illustrates the collaboration between the User, React Client, FastAPI Gateway, PostgreSQL Database, "
        "and WebSocket Alert Dispatcher during authentication and incident triage."
    )
    add_figure_image(
        doc,
        os.path.join(diagrams_dir, "uml_09_collaboration.png"),
        "Figure 4.9: Collaboration diagram for emergency fall detection and response",
        width_inches=5.8
    )

    # 4.3 USER INTERFACE DESIGN USING FIGMA
    doc.add_page_break()
    add_heading_1(doc, "4.3 USER INTERFACE DESIGN USING FIGMA")
    add_body_p(
        doc,
        f"This section presents the user interface design screens and live implemented interfaces of {project_short}. "
        "The interface is architected around role-governed application shells: CareAppShell for clinical staff, "
        "SpaceAppShell for corporate workplace managers, FamilyPortalShell for remote relatives, and SystemAdminShell for administrators. "
        "The interface emphasizes high visual clarity, glanceable room status matrices, and non-blocking alert banners."
    )

    screenshots_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "screenshots")

    # Form Name: Landing page
    add_heading_2(doc, "Form Name: Landing Page")
    add_body_p(
        doc,
        "The landing page introduces the platform, providing an interactive 3D Wi-Fi CSI signal propagation wireframe globe, "
        "channel state telemetry cards (Signal RSSI, Variance, CSI Frame counter), and direct links to role-based authentication. See Figure 4.10."
    )
    lp_img = os.path.join(screenshots_dir, "debug_screen.png")
    if os.path.exists(lp_img):
        add_figure_image(doc, lp_img, "Figure 4.10: Landing page and 3D CSI signal visualizer of Wi-Fi Sense", width_inches=5.6)

    # Form Name: Sign In
    add_heading_2(doc, "Form Name: Sign In")
    add_body_p(
        doc,
        "The sign-in page offers secure credential authentication, role-based contextual routing, password masking, "
        "and quick-fill demo accounts for System Admin, Caregiver, Facility Manager, and Corporate Staff. See Figure 4.11."
    )
    si_img = os.path.join(screenshots_dir, "test_case_1_login.png")
    if os.path.exists(si_img):
        add_figure_image(doc, si_img, "Figure 4.11: Sign-in and authentication page", width_inches=5.6)

    # Form Name: Central Monitoring Dashboard
    add_heading_2(doc, "Form Name: Central Monitoring Dashboard")
    add_body_p(
        doc,
        "The Central Monitoring dashboard displays real-time room occupancy cards across wings, active fall risk badges, "
        "subcarrier amplitude spectrogram streams, and a prominent emergency fall alert banner with one-tap Acknowledge and Resolve actions. See Figure 4.12."
    )
    cm_img = os.path.join(screenshots_dir, "test_case_2_fall_alert.png")
    if os.path.exists(cm_img):
        add_figure_image(doc, cm_img, "Figure 4.12: Central Monitoring dashboard with active emergency fall banner", width_inches=5.6)

    # Form Name: Family Data Sharing Policies
    add_heading_2(doc, "Form Name: Family Data Sharing Policies")
    add_body_p(
        doc,
        "The Family Data Sharing Policies page allows facility managers to configure zero-leak privacy policies, toggling "
        "room presence disclosures, activity details, room suite identifiers, and alert severity thresholds for external family members. See Figure 4.13."
    )
    sp_img = os.path.join(screenshots_dir, "test_case_3_sharing_policy.png")
    if os.path.exists(sp_img):
        add_figure_image(doc, sp_img, "Figure 4.13: Family data sharing and privacy policy configuration", width_inches=5.6)

    # Form Name: Family Care Portal
    add_heading_2(doc, "Form Name: Family Care Portal")
    add_body_p(
        doc,
        "The Family Care Portal allows approved family sponsors to verify resident presence, review safety incident logs, "
        "and manage monthly Care Plan subscriptions without receiving intrusive visual imagery. See Figure 4.14."
    )
    fp_img = os.path.join(screenshots_dir, "test_case_4_family_portal.png")
    if os.path.exists(fp_img):
        add_figure_image(doc, fp_img, "Figure 4.14: Family Care Portal and resident wellness view", width_inches=5.6)

    # Form Name: Corporate Space & Energy Intelligence
    add_heading_2(doc, "Form Name: Corporate Space & Energy Intelligence")
    add_body_p(
        doc,
        "The Corporate Space & Energy Intelligence dashboard renders meeting room occupancy heatmaps, calculates peak utilization rates, "
        "and highlights vacant rooms wasting HVAC power, providing one-click energy mitigation controls. See Figure 4.15."
    )
    cs_img = os.path.join(screenshots_dir, "test_case_5_corporate_space.png")
    if os.path.exists(cs_img):
        add_figure_image(doc, cs_img, "Figure 4.15: Corporate workplace space and energy intelligence dashboard", width_inches=5.6)

    # Form Name: Hardware IoT Device Service Desk
    add_heading_2(doc, "Form Name: Elder Care Sensing Nodes & Service Desk")
    add_body_p(
        doc,
        "The Sensing Nodes page displays the hardware sensor fleet, ESP32 MAC addresses, authentication tokens, assigned room locations, "
        "firmware versions, and hardware fault tracer tickets. See Figure 4.16."
    )
    ds_img = os.path.join(screenshots_dir, "test_case_6_device_service_desk.png")
    if os.path.exists(ds_img):
        add_figure_image(doc, ds_img, "Figure 4.16: Elder Care sensing nodes and hardware service desk", width_inches=5.6)

    # Form Name: Resident Accounts Directory
    add_heading_2(doc, "Form Name: Resident Accounts Directory")
    add_body_p(
        doc,
        "The Resident Accounts page manages resident profiles, assigned room numbers, emergency contact records, "
        "and one-touch emergency inspection actions. See Figure 4.17."
    )
    rd_img = os.path.join(screenshots_dir, "test_case_7_residents_directory.png")
    if os.path.exists(rd_img):
        add_figure_image(doc, rd_img, "Figure 4.17: Resident accounts directory and room assignments", width_inches=5.6)

    # Form Name: Multi-Tenant System Administration Console
    add_heading_2(doc, "Form Name: Multi-Tenant System Administration Console")
    add_body_p(
        doc,
        "The System Administration console provides global tenant oversight across elder care and corporate organizations, "
        "system metrics, telemetry simulation controls, and administrative audit logging. See Figure 4.18."
    )
    sa_img = os.path.join(screenshots_dir, "test_case_8_system_admin_console.png")
    if os.path.exists(sa_img):
        add_figure_image(doc, sa_img, "Figure 4.18: Multi-tenant system administration console", width_inches=5.6)

    # Table 4.1: Application pages declared in the routing module
    add_heading_2(doc, "Application Pages Declared in Routing Module")
    routing_headers = ["Page Component", "Route Path", "Allowed Roles", "Purpose"]
    routing_rows = [
        ["LandingSplash", "/", "Public", "3D CSI wireframe visualization and platform introduction"],
        ["AuthPage / Login", "/login", "Public", "User authentication, JWT issuance, and demo role shortcuts"],
        ["CareAppShell", "/dashboard (care)", "Caregiver, Facility Manager", "Clinical nursing console with live room feed and siren alerts"],
        ["SpaceAppShell", "/dashboard (space)", "Corporate Staff, Facility Admin", "Workplace occupancy matrix and HVAC power mitigation"],
        ["FamilyPortalShell", "/family", "Family Member / Guardian", "Non-visual resident wellness check, safety history, Care Plan"],
        ["DoctorPortalShell", "/doctor", "Doctor / Physician", "Clinical resident review, vital sign telemetry, doctor notes"],
        ["SystemAdminShell", "/admin", "System Administrator", "Global multi-tenant governance, organization setup, audit logs"],
        ["DevicesView", "/devices", "Facility Manager, Admin", "ESP32-S3 sensor fleet inventory, MAC tokens, fault ticketing"],
        ["ResidentsView", "/residents", "Caregiver, Facility Manager", "Resident account directory, room assignments, emergency links"],
        ["SharingPolicyManager", "/policies", "Facility Manager", "Zero-leak family telemetry disclosure and severity filters"],
        ["AlertsView", "/alerts", "Caregiver, Manager, Admin", "Emergency incident history, resolution notes, and audit trail"]
    ]
    add_styled_table(
        doc,
        routing_headers,
        routing_rows,
        [1.6, 1.4, 1.6, 2.0],
        "Table 4.1: Application Pages Declared in Frontend Routing Architecture"
    )

    # 4.4 DATABASE DESIGN
    doc.add_page_break()
    add_heading_1(doc, "4.4 DATABASE DESIGN")
    add_body_p(
        doc,
        "A database is an organized collection of structured data designed to support efficient storage, retrieval, "
        "and relational integrity. In Wi-Fi Sense, the persistence layer manages organizations, buildings, rooms, "
        "sensing nodes, high-frequency CSI telemetry events, activity classifications, and emergency alert logs."
    )

    add_heading_2(doc, "4.4.1 Relational Database Management System (RDBMS)")
    add_body_p(
        doc,
        "PostgreSQL serves as the RDBMS for Wi-Fi Sense. The core relational concepts implemented include:"
    )
    add_bullet_p(doc, "Primary Key: Uniquely identifies each tuple in an entity table (UUID, SERIAL, or BIGSERIAL).", "Primary Key")
    add_bullet_p(doc, "Foreign Key: Enforces referential integrity by linking child entity records to parent tables.", "Foreign Key")
    add_bullet_p(doc, "Entity Integrity: Guarantees that primary key values are non-null and uniquely distinct.", "Entity Integrity")
    add_bullet_p(doc, "Referential Integrity: Ensures that foreign keys reference valid parent rows with cascading deletes where appropriate.", "Referential Integrity")
    add_bullet_p(doc, "Domain Constraints: Restricts column values to valid enumerations (ORGANIZATION_TYPE, ALERT_STATUS_TYPE).", "Domain Constraints")
    add_bullet_p(doc, "Candidate Key: Minimal superkey candidate capable of uniquely distinguishing rows.", "Candidate Key")
    add_bullet_p(doc, "Super Key: Any set of attributes uniquely identifying records within a relational table.", "Super Key")

    add_heading_2(doc, "4.4.2 Normalization")
    add_body_p(
        doc,
        "Normalization is the systematic process of organizing database relations to minimize data redundancy and "
        "prevent insertion, update, and deletion anomalies. The design of Wi-Fi Sense enforces normalization up to "
        "Third Normal Form (3NF). Below is an illustrative walkthrough of the normalization stages:"
    )

    add_heading_3(doc, "First Normal Form (1NF)")
    add_body_p(
        doc,
        "A table is in First Normal Form when each column contains atomic values, each record is uniquely identifiable, "
        "and there are no repeating groups or comma-separated lists."
    )
    add_body_p(doc, "Example (Before 1NF — Unnormalized Table):")
    unnorm_headers = ["Student_ID", "Student_Name", "Courses", "Contact"]
    unnorm_rows = [
        ["101", "Alice", "Math, Science", "9876543210"],
        ["102", "Bob", "English", "8765432109"],
        ["103", "Charlie", "Science, History", "7654321098"]
    ]
    add_styled_table(doc, unnorm_headers, unnorm_rows, [1.2, 1.6, 2.0, 1.6], "Table: Unnormalized Example Relation (Before 1NF)")

    add_body_p(doc, "Conversion to 1NF (Atomic Attributes):")
    norm1_headers = ["Student_ID", "Student_Name", "Course", "Contact"]
    norm1_rows = [
        ["101", "Alice", "Math", "9876543210"],
        ["101", "Alice", "Science", "9876543210"],
        ["102", "Bob", "English", "8765432109"],
        ["103", "Charlie", "Science", "7654321098"],
        ["103", "Charlie", "History", "7654321098"]
    ]
    add_styled_table(doc, norm1_headers, norm1_rows, [1.2, 1.6, 2.0, 1.6], "Table: Relation Converted to First Normal Form (1NF)")

    add_heading_3(doc, "Second Normal Form (2NF)")
    add_body_p(
        doc,
        "A relation is in Second Normal Form (2NF) if it is in 1NF and all non-key attributes are fully functionally "
        "dependent on the complete primary key, removing partial dependencies on composite keys."
    )
    add_body_p(doc, "Conversion to 2NF (Decomposed into Students and Enrollment Tables):")
    s_headers = ["Student_ID", "Student_Name", "Contact"]
    s_rows = [
        ["101", "Alice", "9876543210"],
        ["102", "Bob", "8765432109"],
        ["103", "Charlie", "7654321098"]
    ]
    add_styled_table(doc, s_headers, s_rows, [1.6, 2.4, 2.4], "Table: Students Relation (2NF)")

    e_headers = ["Student_ID", "Course"]
    e_rows = [
        ["101", "Math"],
        ["101", "Science"],
        ["102", "English"],
        ["103", "Science"],
        ["103", "History"]
    ]
    add_styled_table(doc, e_headers, e_rows, [2.5, 3.9], "Table: Course Enrollment Relation (2NF)")

    add_heading_3(doc, "Third Normal Form (3NF)")
    add_body_p(
        doc,
        "A relation is in Third Normal Form (3NF) if it is in 2NF and there are no transitive dependencies: "
        "non-key attributes must not depend on other non-key attributes. In Wi-Fi Sense, rooms reference floor IDs, "
        "floors reference building IDs, and buildings reference organization IDs, ensuring zero transitive updates."
    )

    add_heading_2(doc, "4.4.3 Sanitization")
    add_body_p(
        doc,
        "Data sanitization and input validation techniques protect Wi-Fi Sense against injection attacks and malformed telemetry:"
    )
    add_bullet_p(doc, "Input Validation: Pydantic schemas enforce strict type validation, field length bounds, and regex formats.", "Validation")
    add_bullet_p(doc, "Output Encoding: React automatically encodes dynamic variables within JSX to prevent Cross-Site Scripting (XSS).", "XSS Defense")
    add_bullet_p(doc, "SQL Injection Prevention: SQLModel/SQLAlchemy executes parameterized queries exclusively, avoiding string concatenation.", "SQLi Defense")
    add_bullet_p(doc, "API Validation: Ingestion endpoints validate MAC addresses and node tokens prior to parsing CSI payloads.", "API Security")
    add_bullet_p(doc, "Credential Protection: Passwords are salted and hashed using Argon2/BCrypt; secrets are never logged in responses.", "Credential Protection")

    add_heading_2(doc, "4.4.4 Indexing")
    add_body_p(
        doc,
        "Database indexing significantly accelerates query performance for real-time room monitoring and incident lookups. "
        "PostgreSQL automatically builds B-tree indexes on primary keys and unique constraints. In Wi-Fi Sense, additional indexes "
        "are established on foreign keys and high-frequency query columns:"
    )
    sql_idx = (
        "CREATE INDEX idx_sensing_events_room_timestamp ON sensing_events (room_id, timestamp DESC);\n"
        "CREATE INDEX idx_alerts_room_status ON alerts (room_id, status);\n"
        "CREATE INDEX idx_devices_mac ON devices (mac_address);"
    )
    add_code_block(doc, sql_idx)

    # 4.5 TABLE DESIGN (All 15 tables from schemas.py)
    doc.add_page_break()
    add_all_database_schemas(doc, cfg=cfg)
